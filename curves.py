"""Curves and inflation for the dashboard's Curves page -> docs/curves.js

  Government yields  2Y, 5Y, 10Y, 30Y for the US, Germany, UK, Japan, France,
                     Italy and Canada. History comes from each country's official
                     daily curve where one is free (US Treasury, Bundesbank,
                     Bank of England, Japan's Ministry of Finance, Bank of
                     Canada), lifted by a constant so its level matches the
                     benchmark bond; days the official curve has not reached
                     yet come from CNBC's intraday prices at that market's
                     close, and today from CNBC's quote (about 15 minutes
                     delayed). France and Italy have no free official curve, so
                     they use CNBC alone: weekend-dated bars and one-day spikes
                     are dropped, and each day is re-recorded at the close.
  US curve shape     Treasury's official daily par yield curve, 1M to 30Y
  US inflation       Treasury's official daily real (TIPS) yield curve;
                     breakeven = nominal par yield - real yield, same maturity,
                     and the 5y5y forward breakeven = 2 x 10Y - 5Y

Spreads (2s10s, 5s30s, cross-market) are worked out in the page from the
levels, so every definition sits in one place. History is refetched once a
day; other runs only update the latest values.

python curves.py [--full] [--dump DIR | --cache DIR]
"""
import csv, hashlib, io, json, os, subprocess, sys, time, zipfile
from zoneinfo import ZoneInfo
from datetime import date, datetime, timedelta, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
STORE = os.path.join(HERE, "data", "curves.json")
OUT = os.path.join(HERE, "docs", "curves.js")
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/129.0 Safari/537.36")
KEEP_DAYS = 5 * 366 + 45
COUNTRIES = {   # CNBC symbols for 2Y, 5Y, 10Y, 30Y
    "US": ["US2Y", "US5Y", "US10Y", "US30Y"],
    "DE": ["DE2Y-DE", "DE5Y-DE", "DE10Y-DE", "DE30Y-DE"],
    "UK": ["GB2Y-GB", "GB5Y-GB", "GB10Y-GB", "GB30Y-GB"],
    "JP": ["JP2Y-JP", "JP5Y-JP", "JP10Y-JP", "JP30Y-JP"],
    "FR": ["FR2Y-FR", "FR5Y-FR", "FR10Y-FR", "FR30Y-FR"],
    "IT": ["IT2Y-IT", "IT5Y-IT", "IT10Y-IT", "IT30Y-IT"],
    "CA": ["CA2Y", "CA5Y", "CA10Y", "CA30Y"],
}
TENORS = ["2Y", "5Y", "10Y", "30Y"]
# each market's close, local time, for reading a daily close off CNBC's intraday bars
CLOSE = {"US": ("America/New_York", 17, 0), "DE": ("Europe/Berlin", 17, 30), "FR": ("Europe/Berlin", 17, 30),
         "IT": ("Europe/Berlin", 17, 30), "UK": ("Europe/London", 16, 30), "JP": ("Asia/Tokyo", 15, 0),
         "CA": ("America/Toronto", 16, 0)}
ET = ZoneInfo("America/New_York")
UST_COLS = ["1 Mo", "2 Mo", "3 Mo", "4 Mo", "6 Mo", "1 Yr", "2 Yr", "3 Yr", "5 Yr", "7 Yr", "10 Yr", "20 Yr", "30 Yr"]
REAL_COLS = ["5 YR", "7 YR", "10 YR", "20 YR", "30 YR"]
health = {"errors": [], "requests": 0}
MODE = {"dump": None, "cache": None}


def get(url):
    key = hashlib.sha1(url.encode()).hexdigest()[:16]
    if MODE["cache"]:
        p = os.path.join(MODE["cache"], key)
        if not os.path.exists(p):
            raise IOError("not in cache")
        return open(p, encoding="utf-8").read()
    health["requests"] += 1
    last = ""
    for attempt in range(3):
        p = subprocess.run(["curl", "-sS", "-L", "-m", "60", "--compressed", "-A", UA,
                            "-w", "\n@@%{http_code}", url], capture_output=True)
        body, _, code = p.stdout.decode("utf-8", "replace").rpartition("\n@@")
        if code == "200":
            if MODE["dump"]:
                os.makedirs(MODE["dump"], exist_ok=True)
                open(os.path.join(MODE["dump"], key), "w", encoding="utf-8").write(body)
            return body
        last = f"HTTP {code or 'error'}"
        if code in ("400", "404"):
            break
        time.sleep(2 + 4 * attempt)
    raise IOError(last)


def note(m):
    health["errors"].append(str(m)[:200])


def num(v):
    try:
        return float(str(v).replace("%", "").replace(",", ""))
    except (TypeError, ValueError):
        return None


def cnbc_chart(sym, rng):
    d = json.loads(get(f"https://ts-api.cnbc.com/harmony/app/charts/{rng}.json?symbol={sym}"))
    out = {}
    for b in (d.get("barData") or {}).get("priceBars") or []:
        c, t = num(b.get("close")), str(b.get("tradeTime") or "")
        if c is not None and len(t) >= 8:
            out[f"{t[:4]}-{t[4:6]}-{t[6:8]}"] = c
    return out


def cnbc_quotes(syms):
    out = {}
    for i in range(0, len(syms), 20):
        url = ("https://quote.cnbc.com/quote-html-webservice/restQuote/symbolType/symbol?symbols="
               + "|".join(syms[i:i + 20]) + "&requestMethod=itv&noform=1&partnerId=2&fund=1&exthrs=1&output=json&events=1")
        try:
            for q in json.loads(get(url))["FormattedQuoteResult"]["FormattedQuote"]:
                last, t = num(q.get("last")), q.get("last_time")
                if last is None or not t:
                    continue
                dt = datetime.fromisoformat(t.replace("Z", "+00:00"))
                if dt.date().weekday() < 5:
                    out[q["symbol"]] = (dt.date().isoformat(), last,
                                        dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
        except Exception as e:  # noqa: BLE001
            note(f"CNBC quotes: {e}")
    return out


def treasury(kind, year):
    url = ("https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/"
           f"{year}/all?type={kind}&field_tdr_date_value={year}&page&_format=csv")
    rows = list(csv.DictReader(io.StringIO(get(url))))
    out = {}
    for r in rows:
        try:
            d = datetime.strptime(r["Date"], "%m/%d/%Y").date().isoformat()
        except (KeyError, ValueError):
            continue
        out[d] = {k.strip().upper(): num(v) for k, v in r.items() if k != "Date" and num(v) is not None}
    return out


def cnbc_closes(sym, country):
    """Daily closes read off CNBC's five-day intraday bars (times are New York):
    the last bar at or before the market's close, for days that have finished."""
    d = json.loads(get(f"https://ts-api.cnbc.com/harmony/app/charts/5D.json?symbol={sym}"))
    tzname, hh, mm = CLOSE[country]
    tz = ZoneInfo(tzname)
    best = {}
    for b in (d.get("barData") or {}).get("priceBars") or []:
        c, t = num(b.get("close")), str(b.get("tradeTime") or "")
        if c is None or len(t) < 12:
            continue
        loc = datetime(int(t[:4]), int(t[4:6]), int(t[6:8]), int(t[8:10]), int(t[10:12]), tzinfo=ET).astimezone(tz)
        if loc.weekday() >= 5:
            continue
        close = loc.replace(hour=hh, minute=mm, second=0)
        if loc <= close and (close - loc).total_seconds() <= 90 * 60:
            k = loc.date().isoformat()
            if k not in best or loc > best[k][0]:
                best[k] = (loc, c)
    now = datetime.now(tz)
    return {k: v for k, (t, v) in best.items()
            if k < now.date().isoformat() or now >= t.replace(hour=hh, minute=mm)}


# ---- official daily curves: {tenor: {date: yield}}
def off_bundesbank(start):
    out = {}
    for t, code in (("2Y", "R02XX"), ("5Y", "R05XX"), ("10Y", "R10XX"), ("30Y", "R30XX")):
        txt = get("https://api.statistiken.bundesbank.de/rest/data/BBSIS/D.I.ZST.ZI.EUR.S1311.B.A604."
                  f"{code}.R.A.A._Z._Z.A?format=csv&lang=en&startPeriod={start}")
        out[t] = {}
        for line in txt.splitlines():
            p = line.split(",")
            if len(p) >= 2 and len(p[0]) == 10 and p[0][:2] == "20" and num(p[1]) is not None:
                out[t][p[0]] = num(p[1])
    return out


def off_boc(start):
    ids = {"2Y": "BD.CDN.2YR.DQ.YLD", "5Y": "BD.CDN.5YR.DQ.YLD", "10Y": "BD.CDN.10YR.DQ.YLD", "30Y": "BD.CDN.LONG.DQ.YLD"}
    d = json.loads(get("https://www.bankofcanada.ca/valet/observations/" + ",".join(ids.values())
                       + f"/json?start_date={start}"))
    out = {t: {} for t in ids}
    for o in d.get("observations", []):
        for t, k in ids.items():
            v = num((o.get(k) or {}).get("v"))
            if v is not None:
                out[t][o["d"]] = v
    return out


def off_mof(full):
    urls = (["https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/historical/jgbcme_all.csv"] if full else []) \
        + ["https://www.mof.go.jp/english/policy/jgbs/reference/interest_rate/jgbcme.csv"]
    out = {t: {} for t in TENORS}
    for url in urls:
        rows = list(csv.reader(io.StringIO(get(url))))
        hdr = rows[1]
        for r in rows[2:]:
            try:
                y, m, dd = r[0].split("/")
                k = f"{int(y):04d}-{int(m):02d}-{int(dd):02d}"
            except ValueError:
                continue
            for t in TENORS:
                v = num(r[hdr.index(t)]) if t in hdr else None
                if v is not None:
                    out[t][k] = v
    return out


def off_boe(first):
    """Bank of England nominal zero-coupon (spot) curve; the full archive once, the current month after."""
    import openpyxl
    base = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/"
    out = {t: {} for t in TENORS}
    for name in (["glcnominalddata.zip"] if first else []) + ["latest-yield-curve-data.zip"]:
        p = subprocess.run(["curl", "-sS", "-L", "-m", "180", "-A", UA, base + name], capture_output=True)
        z = zipfile.ZipFile(io.BytesIO(p.stdout))
        for fn in z.namelist():
            if "Nominal daily" not in fn or not ("current month" in fn or "2016" in fn or "2025" in fn):
                continue
            wb = openpyxl.load_workbook(io.BytesIO(z.read(fn)), read_only=True, data_only=True)
            ws = wb[[n for n in wb.sheetnames if n.startswith("4.")][0]]
            cols = None
            for r in ws.iter_rows(values_only=True):
                if r and r[0] == "years:":
                    cols = {t: list(r).index(float(t[:-1])) for t in TENORS if float(t[:-1]) in r}
                elif cols and r and isinstance(r[0], datetime):
                    for t, j in cols.items():
                        if isinstance(r[j], (int, float)):
                            out[t][r[0].date().isoformat()] = round(r[j], 4)
    return out


def despike(dct, jump=0.10, back=0.05):
    """drop a one-day print that jumps and comes straight back"""
    ds = sorted(dct)
    bad = set()
    for i in range(1, len(ds) - 1):
        a, b, c = dct[ds[i - 1]], dct[ds[i]], dct[ds[i + 1]]
        if abs(b - a) > jump and abs(b - c) > jump and (b - a) * (b - c) > 0 and abs(c - a) < back:
            bad.add(ds[i])
    return {d: v for d, v in dct.items() if d not in bad}


def weekdays(dct):
    return {d: v for d, v in dct.items() if date.fromisoformat(d).weekday() < 5}


def main():
    args = sys.argv[1:]
    if "--dump" in args:
        MODE["dump"] = args[args.index("--dump") + 1]
    if "--cache" in args:
        MODE["cache"] = args[args.index("--cache") + 1]
    now = datetime.now(timezone.utc)
    try:
        store = json.load(open(STORE))
    except Exception:  # noqa: BLE001
        store = {}
    full = ("--full" in args or MODE["dump"] or MODE["cache"]
            or store.get("full_date") != now.date().isoformat())
    hist = store.setdefault("hist", {})
    at = store.setdefault("at", {})

    syms = [s for ss in COUNTRIES.values() for s in ss]
    if full:
        for s in syms:
            try:
                h = {}
                weekly, daily = cnbc_chart(s, "5Y"), cnbc_chart(s, "1Y")
                cut = min(daily) if daily else None
                for d, v in weekly.items():
                    if cut is None or d < cut:
                        h[d] = v
                h.update(daily)
                if h:
                    hist[s] = {**hist.get(s, {}), **h}
            except Exception as e:  # noqa: BLE001
                note(f"CNBC chart {s}: {e}")
            if not MODE["cache"]:
                time.sleep(0.25)
    quotes = cnbc_quotes(syms)
    for s, (d, v, t) in quotes.items():
        at[s] = t

    # daily closes off the intraday bars, kept and added to day by day
    closes = store.setdefault("closes", {})
    for c, ss in COUNTRIES.items():
        for s in ss:
            try:
                closes.setdefault(s, {}).update(cnbc_closes(s, c))
            except Exception as e:  # noqa: BLE001
                note(f"CNBC closes {s}: {e}")
            if not MODE["cache"]:
                time.sleep(0.2)

    # official curves, refetched in full once a day
    off = store.setdefault("official", {})
    start = (now.date() - timedelta(days=KEEP_DAYS)).isoformat()
    for c, fn in (("DE", lambda: off_bundesbank(start if full else (now.date() - timedelta(days=20)).isoformat())),
                  ("CA", lambda: off_boc(start if full else (now.date() - timedelta(days=20)).isoformat())),
                  ("JP", lambda: off_mof(full)),
                  ("UK", lambda: off_boe(not store.get("boe_archive")))):
        if not full and c in off and store.get("off_hour", {}).get(c) == now.strftime("%Y-%m-%dT%H"):
            continue
        try:
            got = fn()
            if all(got[t] for t in TENORS):
                for t in TENORS:
                    off.setdefault(c, {}).setdefault(t, {}).update(got[t])
                store.setdefault("off_hour", {})[c] = now.strftime("%Y-%m-%dT%H")
                if c == "UK":
                    store["boe_archive"] = True
            else:
                note(f"official {c}: some tenors empty")
        except Exception as e:  # noqa: BLE001
            note(f"official {c}: {e}")

    # Treasury curves: every year back five on a full run, this year otherwise
    ust = store.setdefault("ust", {})
    real = store.setdefault("real", {})
    years = range(now.year - 5, now.year + 1) if full else [now.year]
    for y in years:
        for kind, dest in (("daily_treasury_yield_curve", ust), ("daily_treasury_real_yield_curve", real)):
            try:
                dest.update(treasury(kind, y))
            except Exception as e:  # noqa: BLE001
                note(f"Treasury {kind} {y}: {e}")

    cutoff = (now.date() - timedelta(days=KEEP_DAYS)).isoformat()
    for s in list(hist):
        hist[s] = {d: v for d, v in hist[s].items() if d >= cutoff}
    for dct in (ust, real):
        for d in [d for d in dct if d < cutoff]:
            del dct[d]
    for s in list(closes):
        closes[s] = {d: v for d, v in closes[s].items() if d >= cutoff}
    for c in off:
        for t in off[c]:
            off[c][t] = {d: v for d, v in off[c][t].items() if d >= cutoff}

    # breakevens from the official curves, same maturity
    be = {}
    for d in sorted(set(ust) & set(real)):
        n, r = ust[d], real[d]
        row = {}
        for t, nk, rk in (("5Y", "5 YR", "5 YR"), ("7Y", "7 YR", "7 YR"), ("10Y", "10 YR", "10 YR"),
                          ("20Y", "20 YR", "20 YR"), ("30Y", "30 YR", "30 YR")):
            if nk in n and rk in r:
                row[t] = round(n[nk] - r[rk], 3)
        if "5Y" in row and "10Y" in row:
            row["5y5y"] = round(2 * row["10Y"] - row["5Y"], 3)
        if row:
            be[d] = row

    def series(dct):
        return sorted(dct.items())

    # the US official curve is Treasury's par curve, already stored for the curve chart
    off["US"] = {t: {d: ust[d][k] for d in ust if k in ust[d]}
                 for t, k in (("2Y", "2 YR"), ("5Y", "5 YR"), ("10Y", "10 YR"), ("30Y", "30 YR"))}
    today = {s: (d, v) for s, (d, v, t) in quotes.items()}
    countries, src = {}, {}
    for c, ss in COUNTRIES.items():
        countries[c] = {}
        for t, s in zip(TENORS, ss):
            cl = closes.get(s, {})
            raw = despike(weekdays(hist.get(s, {})))
            o = (off.get(c) or {}).get(t) or {}
            if len(o) > 40:
                # official history, lifted so its level matches the benchmark bond:
                # the median gap to CNBC over the last 40 shared days
                ref = {**raw, **cl}
                k = sorted(d for d in o if d in ref)[-40:]
                gaps = sorted(ref[d] - o[d] for d in k)
                lift = gaps[len(gaps) // 2] if gaps else 0.0
                ser = {d: round(v + lift, 4) for d, v in o.items()}
                src[c] = "official"
            else:
                ser = dict(raw)
                ser.update(cl)
                src[c] = "cnbc"
            last = max(ser) if ser else ""
            for d, v in cl.items():          # days the official curve has not reached yet
                if d > last:
                    ser[d] = v
            if s in today and date.fromisoformat(today[s][0]).weekday() < 5 and today[s][0] > max(ser or [""]):
                ser[today[s][0]] = today[s][1]
            sr = series(ser)
            countries[c][t] = {"sym": s, "d": [x[0] for x in sr], "v": [x[1] for x in sr], "at": at.get(s)}
    ust_dates = sorted(ust)
    payload = {
        "built": now.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "tenors": TENORS,
        "countries": countries,
        "sources": src,
        "ust": {"cols": UST_COLS, "dates": ust_dates,
                "rows": [[ust[d].get(k.upper()) for k in UST_COLS] for d in ust_dates]},
        "real": {"cols": ["5Y", "7Y", "10Y", "20Y", "30Y"], "dates": sorted(real),
                 "rows": [[real[d].get(k) for k in REAL_COLS] for d in sorted(real)]},
        "be": {"cols": ["5Y", "7Y", "10Y", "20Y", "30Y", "5y5y"], "dates": sorted(be),
               "rows": [[be[d].get(k) for k in ["5Y", "7Y", "10Y", "20Y", "30Y", "5y5y"]] for d in sorted(be)]},
        "health": health,
    }
    missing = [f"{c} {t}" for c in countries for t in TENORS if not countries[c][t]["v"]]
    if missing:
        health["missing"] = missing
    if full:
        store["full_date"] = now.date().isoformat()
    if not MODE["cache"]:
        os.makedirs(os.path.dirname(STORE), exist_ok=True)
        json.dump(store, open(STORE, "w"), separators=(",", ":"))
    out = OUT if not MODE["cache"] else os.path.join(MODE["cache"], "curves.js")
    with open(out, "w") as f:
        f.write("window.CURVES = " + json.dumps(payload, separators=(",", ":")) + ";\n")
    print(f"curves: {health['requests']} requests, missing: {missing or 'none'}, errors: {health['errors'] or 'none'}")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:  # noqa: BLE001
        import traceback
        traceback.print_exc()
        print("curves failed:", e, file=sys.stderr)
    sys.exit(0)
