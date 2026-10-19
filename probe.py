import subprocess, zipfile, io
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0 Safari/537.36"
def getb(url):
    p = subprocess.run(["curl","-sS","-L","-m","120","-A",UA,"-w","\n@@%{http_code}",url],capture_output=True)
    b = p.stdout; i = b.rfind(b"\n@@"); return b[i+3:].decode(), b[:i]
subprocess.run(["pip","install","-q","openpyxl"])
import openpyxl
for url in ["https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/latest-yield-curve-data.zip",
            "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/glcnominalddata.zip"]:
    c, b = getb(url); print(url, c, len(b))
    try:
        z = zipfile.ZipFile(io.BytesIO(b)); print(z.namelist())
        for n in z.namelist():
            if "Nominal" in n or "nominal" in n:
                wb = openpyxl.load_workbook(io.BytesIO(z.read(n)), read_only=True, data_only=True)
                print(n, wb.sheetnames)
                for sh in wb.sheetnames[:6]:
                    ws = wb[sh]; rows = list(ws.iter_rows(min_row=1, max_row=8, values_only=True))
                    print(" sheet", sh, [r[:8] for r in rows])
                    rows = list(ws.iter_rows(values_only=True)); print("  last", [r[:8] for r in rows[-2:]], "nrows", len(rows))
                break
    except Exception as e: print("ERR", e)
